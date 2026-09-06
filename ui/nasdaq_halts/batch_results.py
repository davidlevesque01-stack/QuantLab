from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QAbstractItemView,
    QFileDialog,
    QHBoxLayout,
    QMessageBox,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from ui.nasdaq_halts.results_page import METRIC_DEFINITIONS


DISPLAY_METRIC_HEADERS = [
    "Number of Halt Days\n(trading days)",
    "Average Active HALTs\nper Halt Day\n(HALTs / trading day)",
    "Days Since Last Halt\n(calendar days)",
    "Average Time Between\nHalt Days\n(calendar days)",
    "Sequential Halt Days\nIdentified\n(Yes/No)",
    "Sequential Halt-Day\nBlocks\n(blocks)",
    "Average Sequential\nBlock Length\n(trading days)",
    "Maximum Sequential\nBlock Length\n(trading days)",
    "Halt Days at Close\n(trading days)",
    "HALT on the\nSpecified Day?\n(Yes/No)",
    "HALTs on the\nSpecified Day\n(HALTs)",
]


class BatchResultsPage(QWidget):
    """Display and export File Mode analytical results.

    Date and Ticker remain frozen while the metric columns scroll
    horizontally. Both tables share one synchronized vertical position.
    """

    def __init__(self, on_back) -> None:
        super().__init__()
        self._on_back = on_back
        self._rows: list[dict[str, object]] = []
        self._source_directory = str(Path.home() / "Documents")
        self._syncing_scroll = False
        self._build()

    def _build(self) -> None:
        layout = QVBoxLayout(self)
        layout.setContentsMargins(12, 12, 12, 12)
        layout.setSpacing(8)

        table_layout = QHBoxLayout()
        table_layout.setContentsMargins(0, 0, 0, 0)
        table_layout.setSpacing(0)

        # Frozen context columns.
        self.context_table = QTableWidget(0, 2)
        self.context_table.setHorizontalHeaderLabels(["Date", "Ticker"])
        self.context_table.setEditTriggers(
            QTableWidget.EditTrigger.NoEditTriggers
        )
        self.context_table.setSelectionBehavior(
            QTableWidget.SelectionBehavior.SelectRows
        )
        self.context_table.setSelectionMode(
            QAbstractItemView.SelectionMode.SingleSelection
        )
        self.context_table.verticalHeader().setVisible(False)
        self.context_table.setAlternatingRowColors(True)
        self.context_table.horizontalHeader().setFixedHeight(70)
        self.context_table.setColumnWidth(0, 95)
        self.context_table.setColumnWidth(1, 80)
        self.context_table.setHorizontalScrollBarPolicy(
            Qt.ScrollBarPolicy.ScrollBarAlwaysOff
        )
        self.context_table.setVerticalScrollBarPolicy(
            Qt.ScrollBarPolicy.ScrollBarAlwaysOff
        )
        self.context_table.setFixedWidth(177)

        # Scrollable metrics area.
        self.table = QTableWidget(0, 11)
        self.table.setHorizontalHeaderLabels(DISPLAY_METRIC_HEADERS)
        self.table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.table.setSelectionBehavior(
            QTableWidget.SelectionBehavior.SelectRows
        )
        self.table.setSelectionMode(
            QAbstractItemView.SelectionMode.SingleSelection
        )
        self.table.verticalHeader().setVisible(False)
        self.table.setAlternatingRowColors(True)
        self.table.horizontalHeader().setFixedHeight(70)

        for column in range(11):
            self.table.setColumnWidth(column, 135)

        # Keep Date/Ticker aligned with the metric rows during vertical scroll.
        self.table.verticalScrollBar().valueChanged.connect(
            self._sync_context_scroll
        )
        self.context_table.verticalScrollBar().valueChanged.connect(
            self._sync_metrics_scroll
        )

        # Keep selected row visually aligned on both halves.
        self.table.itemSelectionChanged.connect(
            self._sync_selection_from_metrics
        )
        self.context_table.itemSelectionChanged.connect(
            self._sync_selection_from_context
        )

        table_layout.addWidget(self.context_table)
        table_layout.addWidget(self.table, 1)

        export_xlsx = QPushButton("Export XLSX")
        export_xlsx.setDefault(True)
        export_xlsx.clicked.connect(self._export_xlsx)

        export_csv = QPushButton("Export CSV")
        export_csv.clicked.connect(self._export_csv)

        back = QPushButton("Back")
        back.clicked.connect(self._on_back)

        buttons = QHBoxLayout()
        buttons.addWidget(export_xlsx)
        buttons.addWidget(export_csv)
        buttons.addStretch()
        buttons.addWidget(back)

        layout.addLayout(table_layout)
        layout.addLayout(buttons)

    def _sync_context_scroll(self, value: int) -> None:
        if self._syncing_scroll:
            return
        self._syncing_scroll = True
        self.context_table.verticalScrollBar().setValue(value)
        self._syncing_scroll = False

    def _sync_metrics_scroll(self, value: int) -> None:
        if self._syncing_scroll:
            return
        self._syncing_scroll = True
        self.table.verticalScrollBar().setValue(value)
        self._syncing_scroll = False

    def _sync_selection_from_metrics(self) -> None:
        if self._syncing_scroll:
            return
        row = self.table.currentRow()
        if row >= 0:
            self.context_table.selectRow(row)

    def _sync_selection_from_context(self) -> None:
        if self._syncing_scroll:
            return
        row = self.context_table.currentRow()
        if row >= 0:
            self.table.selectRow(row)

    def set_source_directory(self, directory: str) -> None:
        if directory:
            self._source_directory = directory

    def set_rows(self, rows: list[dict[str, object]]) -> None:
        self._rows = rows

        self.context_table.setRowCount(len(rows))
        self.table.setRowCount(len(rows))

        for row_index, row in enumerate(rows):
            context_values = [row["Date"], row["Ticker"]]

            for column, value in enumerate(context_values):
                item = QTableWidgetItem(self._format_value(value))
                item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                self.context_table.setItem(row_index, column, item)

            metric_values = [
                row[f"Metric {i}"] for i in range(1, 12)
            ]

            for column, value in enumerate(metric_values):
                item = QTableWidgetItem(self._format_value(value))
                item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                self.table.setItem(row_index, column, item)

            # Explicitly keep row heights identical on both table halves.
            row_height = max(
                self.context_table.rowHeight(row_index),
                self.table.rowHeight(row_index),
            )
            self.context_table.setRowHeight(row_index, row_height)
            self.table.setRowHeight(row_index, row_height)

        self.context_table.scrollToTop()
        self.table.scrollToTop()
        self.table.horizontalScrollBar().setValue(0)

    @staticmethod
    def _format_value(value: object) -> str:
        if isinstance(value, float):
            rounded = round(value, 2)
            if rounded.is_integer():
                return str(int(rounded))
            return f"{rounded:.2f}".rstrip("0").rstrip(".")
        return str(value)

    def _default_export_name(self, extension: str) -> str:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return str(
            Path(self._source_directory)
            / f"quantlab_nasdaq_halt_results_{timestamp}.{extension}"
        )

    def _export_xlsx(self) -> None:
        if not self._rows:
            return

        path, _ = QFileDialog.getSaveFileName(
            self,
            "Export results",
            self._default_export_name("xlsx"),
            "Excel Files (*.xlsx)",
        )
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Nasdaq HALT Results"

            headers = ["Date", "Ticker"] + [
                name for _, name in METRIC_DEFINITIONS
            ]
            sheet.append(headers)

            for row in self._rows:
                sheet.append(
                    [row["Date"], row["Ticker"]]
                    + [row[f"Metric {i}"] for i in range(1, 12)]
                )

            # Freeze header row plus Date/Ticker columns in the XLSX too.
            sheet.freeze_panes = "C2"
            sheet.auto_filter.ref = sheet.dimensions

            centered = Alignment(horizontal="center", vertical="center")

            for cell in sheet[1]:
                cell.alignment = centered
                cell.font = Font(bold=True)

            for row_cells in sheet.iter_rows(
                min_row=2,
                max_row=sheet.max_row,
                min_col=1,
                max_col=sheet.max_column,
            ):
                for cell in row_cells:
                    cell.alignment = centered

            for column_cells in sheet.columns:
                width = min(
                    max(len(str(cell.value or "")) for cell in column_cells)
                    + 2,
                    55,
                )
                sheet.column_dimensions[
                    column_cells[0].column_letter
                ].width = width

            workbook.save(path)
        except Exception as exc:
            QMessageBox.critical(self, "Export error", str(exc))
            return

        QMessageBox.information(
            self,
            "Export complete",
            f"Saved:\n{path}",
        )

    def _export_csv(self) -> None:
        if not self._rows:
            return

        path, _ = QFileDialog.getSaveFileName(
            self,
            "Export results",
            self._default_export_name("csv"),
            "CSV Files (*.csv)",
        )
        if not path:
            return
        if not path.lower().endswith(".csv"):
            path += ".csv"

        headers = ["Date", "Ticker"] + [
            name for _, name in METRIC_DEFINITIONS
        ]

        try:
            with open(path, "w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow(headers)
                for row in self._rows:
                    writer.writerow(
                        [row["Date"], row["Ticker"]]
                        + [row[f"Metric {i}"] for i in range(1, 12)]
                    )
        except OSError as exc:
            QMessageBox.critical(self, "Export error", str(exc))
            return

        QMessageBox.information(
            self,
            "Export complete",
            f"Saved:\n{path}",
        )
